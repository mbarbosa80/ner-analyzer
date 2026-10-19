"""2.8.1 — what the 22 Sep 2026 test pass found, fixed and pinned.

A full pass over the 2.8.0 manual, screen by screen against the running
application (Marcelo, 22 Sep 2026: *"agora vc é um testador... verifique toda
a última documentação válida e teste tudo"*), turned up these:

* saving a mapping template under an existing name was a **500** that also
  stopped the import — now it replaces that template;
* the saved templates were **never offered back** — the Import screen now
  applies one;
* the wired-connection scope rule matched the bare word *wired*, so a room
  with only a **USB wired mouse** entered the study as a wired connection;
* the workbook the app exports had **no Matching / Match exceptions sheets**
  — only the command-line tools wrote them;
* a hand-set **photographer** type was dropped by a TEAP re-import unless
  the row carried some other decision;
* a NER erased on the drawing and **restored** on the NERs screen came back
  as a row with no marker;
* a marker **move** was audited as ``element.updated`` with the destination
  only — it is ``ner.moved`` with from and to now;
* the **account events** (created, blocked, reset, deleted) went to an
  audit table no screen read;
* the Demand table's **Ports** column did not sort, the Demand type select
  was cut to "pho", a TER/COW/DIVA reference notice spoke of reach circles
  it does not have, and a dozen strings stayed English in Portuguese.
"""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

import openpyxl
import pymupdf
import pytest

TMP = tempfile.mkdtemp(prefix="ner-281-")
os.environ.setdefault("NER_DATA_DIR", TMP)
os.environ.setdefault("NER_DATABASE_URL", f"sqlite:///{TMP}/t281.sqlite3")

from fastapi.testclient import TestClient                          # noqa: E402
from sqlalchemy import select                                      # noqa: E402

from app import migrate                                            # noqa: E402
from app import services as S                                      # noqa: E402
from app.db import SessionLocal, init_db                           # noqa: E402
from app.domain.scope import WIRED_REQUEST_RX, teap_scope, wired_without_ports  # noqa: E402
from app.main import app                                           # noqa: E402
from app.models import (AuditLog, Drawing, DrawingSpace, Element, MappingTemplate,  # noqa: E402
                        NerAsset, Page, Project, TeapFile, TeapSpace)
from app.sample_data import build_teap                             # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "frontend" / "src"


@pytest.fixture(scope="module")
def client():
    init_db()
    c = TestClient(app)
    from app import auth as A
    with SessionLocal() as db:
        try:
            A.create_user(db, username="tester", password="a-long-test-phrase",
                          display_name="Test User", is_admin=True, actor="tests")
        except A.AuthError:
            pass
    assert c.post("/api/auth/login", json={"username": "tester",
                                           "password": "a-long-test-phrase"}).status_code == 200
    return c


def _pdf(path: str, pages: int = 2) -> str:
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page(width=1000, height=700)
    doc.save(path)
    doc.close()
    return path


def _line(product, qty, ports, code="X", row=1):
    return {"row": row, "product": product, "code": code, "qty": qty, "ports": ports,
            "rule": "A", "explanation": ""}


def _project(name) -> int:
    with SessionLocal() as db:
        p = Project(name=name, venue_code="TST")
        db.add(p)
        db.flush()
        d = Drawing(project_id=p.id, filename="d.pdf", stored_path=_pdf(f"{TMP}/{name}.pdf"),
                    sha256="x" * 64, page_count=2)
        db.add(d)
        db.flush()
        for n in (1, 2):
            db.add(Page(drawing_id=d.id, number=n, width=1000, height=700, metres_per_unit=1.0,
                        scale_label="1u=1m", scale_source="manual", scale_confirmed=True,
                        level="LEVEL 01"))
        db.add(TeapFile(project_id=p.id, filename="teap.xlsx", stored_path=f"{TMP}/t.xlsx",
                        sha256="y" * 64))
        db.commit()
        return p.id


def _room(pid, code, name, x, y, page=1) -> int:
    with SessionLocal() as db:
        p = db.get(Project, pid)
        r = DrawingSpace(project_id=pid, drawing_id=p.drawings[0].id, page=page,
                         space_code=code, name=name, x=x, y=y, confidence=0.9)
        db.add(r)
        db.commit()
        return r.id


def _teap(pid, code, name, ports) -> str:
    with SessionLocal() as db:
        tf = db.execute(select(TeapFile).where(TeapFile.project_id == pid)).scalars().first()
        t = TeapSpace(project_id=pid, teap_file_id=tf.id, space_key=f"TST|{code}",
                      space_code=code, space_name=name, total_ports=ports,
                      lines=[_line("Laptop", ports, ports, "LAP", 1)])
        db.add(t)
        db.commit()
        return t.space_key


def _ner(pid, label, x, y, *, kind="NER_NONCRITICAL", total_ports=48, page=1) -> tuple[int, int]:
    """(asset id, element id)."""
    with SessionLocal() as db:
        p = db.get(Project, pid)
        e = Element(project_id=pid, drawing_id=p.drawings[0].id, page=page,
                    kind=kind, label=label, x=x, y=y, origin="manual",
                    confidence=1.0, confirmed=True, total_ports=total_ports)
        db.add(e)
        db.commit()
        migrate.backfill_assets(db, db.get(Project, pid))
        db.refresh(e)
        e.ner_asset.total_ports = total_ports
        db.commit()
        return e.ner_asset_id, e.id


# ---------------------------------------------------------------------------
# 1. The wired-connection rule
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("product", [
    "Accessories - USB Wired Mouse", "Accessories - Wired Keyboard",
    "Wired Headset", "Display - hardwired mount"])
def test_an_accessory_with_a_cable_is_not_a_wired_connection(product):
    assert not WIRED_REQUEST_RX.search(product)
    lines = [_line(product, 4, 0)]
    assert wired_without_ports(lines) == 0
    assert teap_scope(0, lines) == (False, None)


@pytest.mark.parametrize("product", [
    "Wired Connection - Shared Internet Access (SIA) Wired LAN",
    "Wired Connection - Dedicated plug (RJ45) for RF",
    "Wired Connection - Dedicated Internet Access (DIA)",
    "Dedicated plug RJ-45", "Wired LAN drop", "Wired network outlet"])
def test_a_wired_connection_still_counts(product):
    assert WIRED_REQUEST_RX.search(product)
    lines = [_line(product, 3, 0)]
    assert wired_without_ports(lines) == 3
    in_scope, note = teap_scope(0, lines)
    assert in_scope and "3 wired connection" in note


# ---------------------------------------------------------------------------
# 2. Mapping templates: replaced on the same name, and offered back
# ---------------------------------------------------------------------------

def test_saving_a_template_under_an_existing_name_replaces_it(client):
    r = client.post("/api/projects", json={"name": "tpl", "venue_code": "SYN"})
    pid = r.json()["id"]
    path = build_teap(f"{TMP}/tpl.xlsx")
    with open(path, "rb") as fh:
        up = client.post(f"/api/projects/{pid}/teap",
                         files={"file": ("tpl.xlsx", fh, "application/octet-stream")})
    assert up.status_code == 201, up.text
    teap_id, proposed = up.json()["id"], up.json()["proposed"]
    body = {"sheet": proposed["sheet"], "header_row": proposed["header_row"],
            "mapping": proposed["mapping"], "venue_filter": "SYN",
            "save_as_template": "Tester template"}
    assert client.post(f"/api/projects/{pid}/teap/{teap_id}/import", json=body).status_code == 200
    body["mapping"] = {**proposed["mapping"], "comments": None}
    body["header_row"] = proposed["header_row"]
    # Second save, same name: no 500, one row, the newer mapping.
    r = client.post(f"/api/projects/{pid}/teap/{teap_id}/import", json=body)
    assert r.status_code == 200, r.text
    with SessionLocal() as db:
        rows = db.execute(select(MappingTemplate)
                          .where(MappingTemplate.name == "Tester template")).scalars().all()
        assert len(rows) == 1
        assert rows[0].mapping.get("comments") is None
    names = [x["name"] for x in client.get("/api/mapping-templates").json()]
    assert names.count("Tester template") == 1


def test_the_import_screen_offers_the_saved_templates():
    src = (SRC / "pages" / "panels.tsx").read_text(encoding="utf-8")
    assert "api.templates()" in src
    assert 'Apply a saved mapping template' in src
    assert "columns not in this sheet, left as they were" in src
    api = (SRC / "lib" / "api.ts").read_text(encoding="utf-8")
    assert "templates: () => req<any[]>(`/mapping-templates`)" in api


# ---------------------------------------------------------------------------
# 3. The workbook carries the matching
# ---------------------------------------------------------------------------

def test_the_exported_workbook_has_the_matching_sheets(client):
    pid = _project("wbmatch")
    _room(pid, "OVR.001", "OVR Office", 100, 100)
    _room(pid, "OVR.001", "OVR Office", 100, 100, page=2)     # the same code on two pages
    _teap(pid, "OVR.001", "OVR Office", 10)
    _teap(pid, "GHO.099", "Ghost room", 4)                     # nothing on the drawing
    _ner(pid, "NER 1", 150, 120)
    assert client.post(f"/api/projects/{pid}/matching/run").status_code == 200
    assert client.post(f"/api/projects/{pid}/analyse").status_code == 200
    r = client.post(f"/api/projects/{pid}/reports/xlsx")
    assert r.status_code == 200, r.text
    from app.config import settings          # the data dir of whichever module loaded first
    wb = openpyxl.load_workbook(Path(settings.exports_dir) / f"project_{pid}" / r.json()["file"])
    assert "Matching" in wb.sheetnames and "Match exceptions" in wb.sheetnames
    ws = wb["Matching"]
    head = [c.value for c in ws[1]]
    assert head[:6] == ["TEAP code", "TEAP name", "Drawing code", "Drawing name", "Page", "Method"]
    assert "Confirmed by" in head
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["OVR.001"][2] == "OVR.001" and rows["OVR.001"][5] != "none"
    assert rows["GHO.099"][5] == "none"
    ex = list(wb["Match exceptions"].iter_rows(min_row=2, values_only=True))
    kinds = {e[0] for e in ex}
    assert "TEAP not found on drawings" in kinds
    assert any(e[0] == "TEAP not found on drawings" and e[1] == "GHO.099" for e in ex)
    # every sheet the manual lists, in the app's own export ("Wireless points"
    # and "TEAP spaces with no ports" appear only when there are any)
    for sheet in ["Summary", "Coverage figures", "TEAP and coverage", "NERs and ports",
                  "Exceptions", "Matching", "Match exceptions", "Scales"]:
        assert sheet in wb.sheetnames, sheet


# ---------------------------------------------------------------------------
# 4. A hand-set photographer type survives a re-import on its own
# ---------------------------------------------------------------------------

def test_a_hand_set_photographer_type_is_carried_across_a_reimport():
    pid = _project("photo-carry")
    key = _teap(pid, "SEC.101", "Security Waiting", 3)
    with SessionLocal() as db:
        t = db.execute(select(TeapSpace).where(TeapSpace.space_key == key)).scalar_one()
        t.demand_type = "photographer"
        db.commit()
        p = db.get(Project, pid)
        carried = S._lift_decisions(db, p)
    assert key in carried and carried[key]["demand_type"] == "photographer"
    # …and a row with nothing decided at all is not carried
    key2 = _teap(pid, "GEN.001", "General room", 2)
    with SessionLocal() as db:
        assert key2 not in S._lift_decisions(db, db.get(Project, pid))


# ---------------------------------------------------------------------------
# 5. Restore redraws the marker the viewer erased
# ---------------------------------------------------------------------------

def test_restore_brings_the_marker_back_after_the_viewer_erased_it(client):
    # Since 2.18.1 the viewer's plain Delete removes the NER and keeps its
    # markers hidden; the marker-only act (``scope=marker``) is the one that
    # erases the drawing and leaves Restore to redraw it — tested here.
    pid = _project("restore-marker")
    asset_id, elem_id = _ner(pid, "NER 9", 400, 300, kind="NER_CRITICAL")
    assert client.delete(f"/api/projects/{pid}/elements/{elem_id}?scope=marker").status_code == 204
    with SessionLocal() as db:
        a = db.get(NerAsset, asset_id)
        assert a.deleted_at is not None and not list(a.occurrences)
    r = client.post(f"/api/projects/{pid}/ners/{asset_id}/restore")
    assert r.status_code == 200, r.text
    with SessionLocal() as db:
        a = db.get(NerAsset, asset_id)
        assert a.deleted_at is None
        occ = list(a.occurrences)
        assert len(occ) == 1
        m = occ[0]
        assert (m.page, round(m.x), round(m.y), m.kind, m.label) == (1, 400, 300, "NER_CRITICAL", "NER 9")
        assert m.origin == "manual" and m.link_status == "confirmed"
        log = db.execute(select(AuditLog).where(AuditLog.project_id == pid,
                                                 AuditLog.action == "ner.restored")).scalars().all()
        assert log and log[-1].payload.get("marker_redrawn") == m.id
    # the viewer lists it again
    labels = [e["label"] for e in client.get(f"/api/projects/{pid}/elements?page=1").json()]
    assert "NER 9" in labels


# ---------------------------------------------------------------------------
# 6. A move is a move in the audit log
# ---------------------------------------------------------------------------

def test_moving_a_marker_is_audited_as_ner_moved_with_from_and_to(client):
    pid = _project("move-audit")
    asset_id, elem_id = _ner(pid, "NER 3", 100, 100)
    r = client.patch(f"/api/projects/{pid}/elements/{elem_id}", json={"x": 160, "y": 140})
    assert r.status_code == 200
    with SessionLocal() as db:
        moved = db.execute(select(AuditLog).where(AuditLog.project_id == pid,
                                                   AuditLog.action == "ner.moved")).scalars().all()
        assert len(moved) == 1
        pl = moved[0].payload
        assert pl["from"] == {"page": 1, "x": 100.0, "y": 100.0}
        assert pl["to"] == {"page": 1, "x": 160.0, "y": 140.0}
        assert pl["element_id"] == elem_id and moved[0].entity_id == asset_id
        # no "element.updated" for a pure move
        upd = db.execute(select(AuditLog).where(AuditLog.project_id == pid,
                                                 AuditLog.action == "element.updated")).scalars().all()
        assert upd == []
        e = db.get(Element, elem_id)
        assert e.origin == "manual"       # unchanged: a move is not a new origin
    # a TER marker moves as an element
    with SessionLocal() as db:
        p = db.get(Project, pid)
        t = Element(project_id=pid, drawing_id=p.drawings[0].id, page=1, kind="TER",
                    label="TER", x=10, y=10, origin="extracted", confidence=0.8)
        db.add(t)
        db.commit()
        tid = t.id
    assert client.patch(f"/api/projects/{pid}/elements/{tid}", json={"x": 20}).status_code == 200
    with SessionLocal() as db:
        em = db.execute(select(AuditLog).where(AuditLog.project_id == pid,
                                                AuditLog.action == "element.moved")).scalars().all()
        assert len(em) == 1 and em[0].payload["to"]["x"] == 20.0
        assert db.get(Element, tid).origin == "extracted"


# ---------------------------------------------------------------------------
# 7. Account events can be read
# ---------------------------------------------------------------------------

def test_account_events_are_listed_for_administrators(client):
    r = client.post("/api/admin/users", json={"username": "eventful", "password": "a-long-enough-one",
                                              "display_name": "Ev", "is_admin": False})
    assert r.status_code == 201, r.text
    uid = r.json()["id"]
    assert client.post(f"/api/admin/users/{uid}/active", json={"value": False}).status_code == 200
    log = client.get("/api/admin/audit")
    assert log.status_code == 200
    actions = [a["action"] for a in log.json()]
    assert "user.created" in actions and "user.disabled" in actions
    # 4.2.0: the installation's own clean-ups (files of deleted projects) are
    # listed beside the account events — both belong to no project.
    # 4.2.1: and a deleted project's own line (its audit log goes with it).
    assert all(a["entity"] in ("user", "installation", "project") for a in log.json())
    # not for a plain account
    other = TestClient(app)
    with SessionLocal() as db:
        from app import auth as A
        try:
            A.create_user(db, username="plain", password="a-plain-long-phrase", actor="tests")
        except A.AuthError:
            pass
    assert other.post("/api/auth/login", json={"username": "plain",
                                                "password": "a-plain-long-phrase"}).status_code == 200
    assert other.get("/api/admin/audit").status_code == 403
    src = (SRC / "pages" / "Access.tsx").read_text(encoding="utf-8")
    assert 'api.accountAudit()' in src and 't("Account events")' in src


# ---------------------------------------------------------------------------
# 8. The interface details the pass caught
# ---------------------------------------------------------------------------

def test_the_demand_ports_column_sorts_by_the_row_field():
    src = (SRC / "pages" / "panels.tsx").read_text(encoding="utf-8")
    assert '<sd.Th k="ports" get={(r) => r.total_ports}>' in src
    assert 'style={{ minWidth: 170 }}' in src            # the demand-type select


def test_the_audit_payload_column_sorts_as_text():
    src = (SRC / "pages" / "Access.tsx").read_text(encoding="utf-8")
    assert 'get={(a: any) => JSON.stringify(a.payload || {})}' in src


def test_a_non_ner_reference_notice_promises_no_circles():
    src = (SRC / "pages" / "Viewer.tsx").read_text(encoding="utf-8")
    assert "drawn only, no reach circles; nothing is counted" in src
    assert "NER_REFERENCE_KINDS.has(kind)\n          ? t(\"reference placed on this sheet — drawn only, with its reach circles" in src


def test_the_job_bar_does_not_stutter_when_the_stage_is_the_state():
    src = (SRC / "components" / "JobBar.tsx").read_text(encoding="utf-8")
    assert "job.stage.toLowerCase() !== job.state" in src
    assert 't("Finished")' in src and 't("Working")' in src


def test_the_open_project_is_remembered_across_a_reload():
    src = (SRC / "App.tsx").read_text(encoding="utf-8")
    assert 'localStorage.getItem("ner.project")' in src
    assert 'localStorage.setItem("ner.project", String(projectId))' in src
    assert 'localStorage.removeItem("ner.project")' in src


def test_no_raw_english_text_nodes_are_left_on_the_screens():
    """The interface strings the pass found untranslated now go through t()."""
    pt = (SRC / "lib" / "strings.pt.ts").read_text(encoding="utf-8")
    for key in ["Sign out", "Change password", "Reset password", "Delete permanently",
                "Import TEAP", "Reload venue list", "Venue filter", "import every venue",
                "Macros are never executed. Formulas are read as their cached values.",
                "Create account and sign in", "What this does not protect", "metres", "feet",
                "both", "you", "Page number this shaft connects to", "Vertical length in metres",
                "Serving NER", "Method"]:
        assert f'"{key}":' in pt, key
    app_src = (SRC / "App.tsx").read_text(encoding="utf-8")
    assert '{t("Sign out")}' in app_src
    access = (SRC / "pages" / "Access.tsx").read_text(encoding="utf-8")
    for raw in [">\n            Change password\n", ">\n                          Reset password\n",
                ">\n                  Delete permanently\n", "What this does not protect\n      </h4>"]:
        assert raw not in access, raw
    panels = (SRC / "pages" / "panels.tsx").read_text(encoding="utf-8")
    assert "\n                Import TEAP\n" not in panels
    assert '<option value="m">{t("metres")}</option>' in panels
