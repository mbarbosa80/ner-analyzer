"""The criticality matrix: which room types require a critical NER.

This is a **second reference document**, separate from the TEAP and separate
from the drawing, and the distinction took a while to establish so it is worth
writing down.

* The **TEAP** is per venue. It says *this room needs N ports of product X*.
  It has no criticality column at all — 21 columns on the event's venue sheets,
  none of them criticality. That was verified against the real files after an
  earlier guess that the importer was failing to read one.
* The **matrix** is programme-wide. It says *this room **type** requires a
  critical NER, or a non-critical one*, keyed by Space Code.

They join on Space Code, and the join is partial in a way that matters:

    HND, 112 TEAP spaces
        exact space code        23   20.5%
        + parent code          +17   15.2%   ->  36%
        + fuzzy name            +2    1.8%   ->  38%
        no rule at all          70   62.5%

Six rooms in ten have no rule, and they are not obscure ones — Field of Play,
Athlete Medical Station, Central Technical Area, Power Generator. **That is a
gap in the source document, not in this code**, and the honest thing is to
report it rather than paper over it with a default. A room with no stated
criticality stays unknown; it never becomes "non-critical" by omission, because
non-critical is a permission (any room may serve it) and permissions must be
granted, not assumed.

Three ways a rule reaches a space, and they are never conflated:

``stated``
    The matrix names this exact space code. Nothing inferred.
``inherited``
    The matrix names an ancestor code — ``PRS.001.03`` inheriting from
    ``PRS.001``. Defensible, because the hierarchy is real in the taxonomy,
    but the register shows where it came from so a reviewer can disagree with
    a specific case rather than the whole rule.
``proposed``
    Only the room *name* is similar. Never applied automatically: getting this
    wrong promotes or demotes a room, and it recovers less than 2% anyway.
    It becomes an alias for a human to approve.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Iterable, Optional, Sequence

from .teap_import import detect_criticality_columns, is_ticked

CRITICAL = "critical"
NON_CRITICAL = "non-critical"

#: Segments that are placeholders rather than real levels of the hierarchy.
#: ``PRS.DUM.099`` is "some photo position", not the 99th child of DUM.
PLACEHOLDER_SEGMENTS = {"DUM", "DUMMY", "TBC", "TBD", "NA", "XXX"}

#: Below this the name similarity is noise. Above it the pair is *proposed*,
#: never applied — see the module docstring.
NAME_PROPOSAL_SCORE = 88


@dataclass(frozen=True)
class Verdict:
    """What the matrix has to say about one space."""

    criticality: Optional[str]
    provenance: Optional[str]          # stated | inherited | proposed | None
    matched_key: Optional[str] = None
    score: Optional[float] = None

    @property
    def applies(self) -> bool:
        """Should this be written onto the space without asking?"""
        return self.provenance in ("stated", "inherited")

    def explain(self) -> str:
        if self.provenance == "stated":
            return f"{self.criticality} — stated by the matrix for {self.matched_key}"
        if self.provenance == "inherited":
            return (f"{self.criticality} — inherited from {self.matched_key}, the "
                    f"parent space code; the matrix does not name this space directly")
        if self.provenance == "proposed":
            return (f"possibly {self.criticality} — the matrix row “{self.matched_key}” "
                    f"has a similar name ({self.score:.0f}% match). Not applied: "
                    f"confirm it as an alias first")
        return ("not stated in the criticality matrix — neither this space code nor "
                "any parent code appears in it")


@dataclass
class MatrixRow:
    code: Optional[str]
    name: str
    criticality: Optional[str]
    conflict: bool = False


@dataclass
class CriticalityMatrix:
    """The loaded matrix, ready to answer questions about spaces."""

    by_code: dict[str, str] = field(default_factory=dict)
    by_name: dict[str, str] = field(default_factory=dict)
    rows: list[MatrixRow] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)
    blank: list[str] = field(default_factory=list)

    # -- lookup ------------------------------------------------------------

    def lookup(self, space_code: Optional[str], space_name: Optional[str] = None
               ) -> Verdict:
        code = (space_code or "").strip().upper()
        if code and not code.startswith("#") and code in self.by_code:
            return Verdict(self.by_code[code], "stated", code)

        for ancestor in ancestors(code):
            if ancestor in self.by_code:
                return Verdict(self.by_code[ancestor], "inherited", ancestor)

        name = (space_name or "").strip().lower()
        if name and self.by_name:
            best, score = _closest(name, self.by_name)
            if best is not None and score >= NAME_PROPOSAL_SCORE:
                return Verdict(self.by_name[best], "proposed", best, score)

        return Verdict(None, None)

    # -- reporting ---------------------------------------------------------

    def coverage(self, spaces: Iterable[tuple[Optional[str], Optional[str]]]) -> dict:
        """How much of a venue this matrix actually reaches.

        The number nobody wants to see but everybody needs: what fraction of the
        rooms in this venue the programme has written a rule for.
        """
        tally = {"stated": 0, "inherited": 0, "proposed": 0, "unknown": 0}
        gaps: list[dict] = []
        proposals: list[dict] = []
        total = 0
        for code, name in spaces:
            total += 1
            v = self.lookup(code, name)
            tally[v.provenance or "unknown"] += 1
            if v.provenance is None:
                gaps.append({"space_code": code, "space_name": name})
            elif v.provenance == "proposed":
                proposals.append({"space_code": code, "space_name": name,
                                  "matrix_row": v.matched_key,
                                  "criticality": v.criticality,
                                  "score": round(v.score or 0, 1)})
        applied = tally["stated"] + tally["inherited"]
        return {
            "total_spaces": total,
            "applied": applied,
            "applied_pct": round(100 * applied / total, 1) if total else 0.0,
            "by_provenance": tally,
            "rule_missing": len(gaps),
            "rule_missing_pct": round(100 * tally["unknown"] / total, 1) if total else 0.0,
            "gaps": gaps,
            "alias_proposals": proposals,
            "note": (
                "Spaces with no rule are left unknown on purpose. Treating a "
                "missing rule as non-critical would grant a permission the "
                "matrix never gave, and non-critical is a permission: it lets "
                "any room serve the space."),
        }


def ancestors(code: str) -> list[str]:
    """Parent codes, nearest first. ``PRS.001.03`` -> ``PRS.001`` -> ``PRS``.

    Placeholder segments are dropped rather than treated as a level, so
    ``PRS.DUM.099`` walks straight to ``PRS`` instead of inventing ``PRS.DUM``.
    """
    if not code or code.startswith("#"):
        return []
    parts = [p for p in code.split(".") if p and p.upper() not in PLACEHOLDER_SEGMENTS]
    # A trailing "-001" style suffix is an instance number, not a level.
    if parts:
        parts[-1] = re.sub(r"-\d+$", "", parts[-1])
    out = []
    while len(parts) > 1:
        parts = parts[:-1]
        out.append(".".join(parts))
    return out


def _closest(name: str, table: dict[str, str]) -> tuple[Optional[str], float]:
    try:
        from rapidfuzz import fuzz
    except ImportError:                                    # pragma: no cover
        return None, 0.0
    best, score = None, 0.0
    for candidate in table:
        s = fuzz.token_set_ratio(name, candidate)
        if s > score:
            best, score = candidate, s
    return best, score


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def build_matrix(headers: Sequence[Any], data_rows: Iterable[Sequence[Any]],
                 *, code_header: str = "Space Code",
                 name_header: Optional[str] = None) -> CriticalityMatrix:
    """Turn a header row plus data rows into a matrix.

    Split out from file reading so the rules can be tested without a workbook.
    """
    head = [("" if c is None else str(c).strip()) for c in headers]
    pair = detect_criticality_columns(head)
    if not (pair["critical_flag"] or pair["non_critical_flag"]):
        raise ValueError(
            "this sheet has no criticality columns. Expected a pair like "
            "'Critical NER' and 'Non-Critical NER', or a single criticality column")

    lower = [h.lower() for h in head]

    def find(target: str, fallbacks: Sequence[str] = ()) -> Optional[int]:
        for want in (target, *fallbacks):
            w = want.lower()
            for i, h in enumerate(lower):
                if h == w:
                    return i
            for i, h in enumerate(lower):
                if w in h:
                    return i
        return None

    ci = head.index(pair["critical_flag"]) if pair["critical_flag"] else None
    ni = head.index(pair["non_critical_flag"]) if pair["non_critical_flag"] else None
    code_i = find(code_header, ("space code", "room code"))
    name_i = find(name_header or "target room", ("space name", "room", "space"))

    m = CriticalityMatrix()
    for row in data_rows:
        cells = list(row or [])
        if not any(c is not None and str(c).strip() for c in cells):
            continue

        def cell(idx: Optional[int]) -> Any:
            return cells[idx] if idx is not None and idx < len(cells) else None

        code = ("" if cell(code_i) is None else str(cell(code_i)).strip())
        name = ("" if cell(name_i) is None else str(cell(name_i)).strip())
        crit = is_ticked(cell(ci)) if ci is not None else None
        non = is_ticked(cell(ni)) if ni is not None else None

        if crit and non:
            m.conflicts.append(name or code)
            m.rows.append(MatrixRow(code or None, name, None, conflict=True))
            continue
        value = CRITICAL if crit else (NON_CRITICAL if non else None)
        m.rows.append(MatrixRow(code or None, name, value))
        if value is None:
            m.blank.append(name or code)
            continue
        if code and not code.startswith("#"):
            m.by_code[code.upper()] = value
        elif name:
            m.by_name[name.lower()] = value
    return m


def load_matrix_workbook(path: str, *, sheet: Optional[str] = None,
                         header_row: int = 1) -> CriticalityMatrix:
    """Read the matrix from an .xlsx file."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if len(rows) < header_row:
        raise ValueError(f"the sheet has fewer than {header_row} rows")
    return build_matrix(rows[header_row - 1], rows[header_row:])
