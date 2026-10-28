"""2.18.0 — the "NER Standardized Naming Structure" v2.0.

Marcelo, 9 Oct 2026, with *NER Standardized Naming Structure - v2.0.xlsx*:

* non-critical racks are numbered from ``NER 20.1`` (v1.1 said 21);
* every rack has a name of its own, ``NERc NN.k`` — and his form of it is
  *"não usar espaços e sim NERc_01.1, isso serve pra tudo"*;
* *"leva o nome completo para o excel"* — the venue code in front
  (``VHH_NERc_01.1``), on the *NERs and ports* sheet.

Nothing already named is renamed; the register reads the old spellings and the
new ones as one series when it hands out the next number.
"""

from __future__ import annotations

import os
import tempfile

import pytest

TMP = tempfile.mkdtemp(prefix="ner-2180-")
os.environ.setdefault("NER_DATA_DIR", TMP)
os.environ.setdefault("NER_DATABASE_URL", f"sqlite:///{TMP}/t2180.sqlite3")

from app import ner_register as R                                  # noqa: E402
from app.db import SessionLocal, init_db                           # noqa: E402
from app.domain.reports import ner_full_name                       # noqa: E402
from app.models import Project                                     # noqa: E402
from tests.test_wireless_tie_250 import client  # noqa: E402,F401


@pytest.fixture(scope="module", autouse=True)
def _db():
    init_db()


# ---------------------------------------------------------------------------
# The form and the series
# ---------------------------------------------------------------------------

def test_the_app_writes_the_v2_form_with_an_underscore():
    assert R.sequential_name(1, True) == "NERc_01.1"
    assert R.sequential_name(20, False) == "NER_20.1"
    assert R.sequential_name(107, True) == "NERc_107.1"


def test_non_critical_starts_at_20_as_the_v2_sheet_numbers_them():
    assert R.SEQUENCE_START == {True: 1, False: 20}
    assert R.next_sequential_name([], False) == "NER_20.1"
    assert R.next_sequential_name([], True) == "NERc_01.1"
    assert R.next_sequential_name(["NER_20.1"], False) == "NER_21.1"


@pytest.mark.parametrize("name,critical,number", [
    ("NERc_01.1", True, 1),        # the v2.0 form
    ("NER_20.1", False, 20),
    ("NERC_03.2", True, 3),        # case and a second rack
    ("NERc 02.1", True, 2),        # the sheet's own spelling
    ("VHH_NERc_01.1", True, 1),    # the full name
    ("LSH_NERc 02.1", True, 2),
    ("NERc01", True, 1),           # 2.14–2.17 names still in the projects
    ("NERc 01", True, 1),
    ("NER-21", False, 21),
    ("ner 3", False, 3),
])
def test_every_spelling_of_a_sequential_name_is_read(name, critical, number):
    m = R.SEQUENTIAL_NAME_RX.match(name)
    assert m, name
    assert bool(m.group(1)) is critical and int(m.group(2)) == number


@pytest.mark.parametrize("name", ["Critical NER 12", "VEE-NER-01", "My rack", "NER", "TEC.010-003"])
def test_a_hand_name_is_not_a_sequential_name(name):
    assert R.SEQUENTIAL_NAME_RX.match(name) is None


def test_old_new_and_prefixed_spellings_are_one_series():
    names = ["NERc01", "NERc 02", "VHH_NERc_03.1", "NERc_04.2", "NER21", "NER_22.1", "My rack"]
    assert R.next_sequential_name(names, True) == "NERc_05.1"
    assert R.next_sequential_name(names, False) == "NER_23.1"
    # the collision check sees through the spelling: "NERc 05.1" is "NERc_05.1"
    assert R.next_sequential_name(names + ["NERc 05.1"], True) == "NERc_06.1"


# ---------------------------------------------------------------------------
# The full name
# ---------------------------------------------------------------------------

def test_the_full_name_puts_the_venue_code_in_front_once():
    assert ner_full_name("VHH", "NERc_01.1") == "VHH_NERc_01.1"
    assert ner_full_name("vhh", "NER_20.1") == "VHH_NER_20.1"
    assert ner_full_name("VHH", "VHH_NERc_01.1") == "VHH_NERc_01.1", "never twice"
    assert ner_full_name("VHH", "LSH_NERc_02.1") == "LSH_NERc_02.1", "a code already there is kept"
    assert ner_full_name(None, "NERc_01.1") == "NERc_01.1"
    assert ner_full_name("", "NERc_01.1") == "NERc_01.1"
    assert ner_full_name("VHH", "My rack") == "VHH_My rack", "a hand name gets the code too"
    assert ner_full_name("VHH", "") == ""


def test_the_ners_sheet_has_a_full_name_column_with_the_venue_code(client):
    import openpyxl

    from app.config import settings
    from tests.test_wireless_tie_250 import _analyse, _match, _ner, _project, _room, _teap, OFFICE
    pid = _project("fullname-1")                       # venue_code TST
    _room(pid, "OFF.070", "Office", 100, 100)
    _teap(pid, "OFF.070", "Office", OFFICE)
    _ner(pid, "NERc_01.1", 100, 120, total_ports=24)
    _match(client, pid)
    _analyse(pid)
    r = client.post(f"/api/projects/{pid}/reports/xlsx?background=false")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    ws = wb["NERs and ports"]
    head = [c.value for c in ws[1]]
    # 3.0.0: Full name is column B, the space code C, then the old order
    assert head[:5] == ["NER", "Full name", "Space code", "Page", "Critical"]
    row = [c.value for c in ws[2]]
    assert row[0] == "NERc_01.1" and row[1] == "TST_NERc_01.1"
    assert ws.auto_filter.ref.startswith("A1:N")
    # the TOTAL row is still there, and leaves the name column alone
    totals = [r for r in ws.iter_rows(values_only=True) if r and r[0] == "TOTAL"]
    assert len(totals) == 1 and totals[0][1] is None


def test_a_project_without_a_venue_code_gets_the_name_alone(client):
    import openpyxl

    from app.config import settings
    from tests.test_wireless_tie_250 import _analyse, _match, _ner, _project, _room, _teap, OFFICE
    pid = _project("fullname-2")
    with SessionLocal() as db:
        p = db.get(Project, pid)
        p.venue_code = ""
        db.commit()
    _room(pid, "OFF.071", "Office", 100, 100)
    _teap(pid, "OFF.071", "Office", OFFICE)
    _ner(pid, "NER_20.1", 100, 120, total_ports=24)
    _match(client, pid)
    _analyse(pid)
    r = client.post(f"/api/projects/{pid}/reports/xlsx?background=false")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    ws = wb["NERs and ports"]
    assert [c.value for c in ws[1]][1] == "Full name"
    assert [c.value for c in ws[2]][1] == "NER_20.1"
