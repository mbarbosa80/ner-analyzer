"""3.0.0 — a room drawn on several sheets, the NERs sheet, the sheet order,
and an unconfirmed match that is served.

Marcelo, 14 Oct 2026, on the ALA export:

* "quando uma NER aparece em mais de uma página, o PDF traz a NER posicionada
  em local errado … A NERc_04.1 não aparece na página 6 e na página 4 ela está
  errada." Two defects behind it: the exporter drew every room once, at its
  canonical point, and a drag of the room's marker on any sheet copied that
  sheet's x/y into the canonical point while the canonical *page* stayed — so
  NERc_04.1 stood at page 6's coordinates on page 4. The engine measured from
  that wrong point too.
* the notes written on the NERs screen go to *NERs and ports*; Full name moves
  to B, a *Space code* (AAM / VDB) is C.
* the sheet order and the tab colours he gave.
* a TEAP space the matcher placed on a room, with a NER measured, is served by
  the nearest NER even before the match is confirmed.
"""

from __future__ import annotations

import os
import tempfile

import pymupdf
import pytest

TMP = tempfile.mkdtemp(prefix="ner-2200-")
os.environ.setdefault("NER_DATA_DIR", TMP)
os.environ.setdefault("NER_DATABASE_URL", f"sqlite:///{TMP}/t2200.sqlite3")

from fastapi.testclient import TestClient                          # noqa: E402
from sqlalchemy import select                                      # noqa: E402

from app import migrate                                            # noqa: E402
from app import services as S                                      # noqa: E402
from app.db import SessionLocal, init_db                           # noqa: E402
from app.domain.coverage import NER, compute_coverage              # noqa: E402
from app.domain.distance import DemandType, DistanceRules, Status  # noqa: E402
from app.domain.geometry import Area, Point                        # noqa: E402
from app.domain.units import Scale                                 # noqa: E402
from app.main import app                                           # noqa: E402
from app.models import (Drawing, DrawingSpace, Element, NerAsset, Page,  # noqa: E402
                        Project, TeapFile, TeapSpace)


@pytest.fixture(scope="module")
def client():
    init_db()
    c = TestClient(app)
    from app import auth as A
    with SessionLocal() as db:
        try:
            A.create_user(db, username="tester", password="a-long-test-phrase",
                          display_name="Test User", is_admin=True, actor="tests")
        except A.AuthError:
            pass
    assert c.post("/api/auth/login", json={"username": "tester",
                                           "password": "a-long-test-phrase"}).status_code == 200
    return c


def _pdf(path: str, pages: int = 2) -> str:
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page(width=1000, height=700)
    doc.save(path)
    doc.close()
    return path


def _project(name: str) -> tuple[int, int]:
    """Two sheets of the same level, both 1 unit = 1 m."""
    with SessionLocal() as db:
        p = Project(name=name, venue_code="ALA")
        db.add(p)
        db.flush()
        d = Drawing(project_id=p.id, filename="d.pdf", stored_path=_pdf(f"{TMP}/{name}.pdf"),
                    sha256="x" * 64, page_count=2)
        db.add(d)
        db.flush()
        for n in (1, 2):
            db.add(Page(drawing_id=d.id, number=n, width=1000, height=700, metres_per_unit=1.0,
                        scale_label="1u=1m", scale_source="manual", scale_confirmed=True,
                        level="LEVEL 01"))
        db.add(TeapFile(project_id=p.id, filename="teap.xlsx", stored_path=f"{TMP}/t.xlsx",
                        sha256="y" * 64))
        db.commit()
        return p.id, d.id


def _room_on_two_sheets(pid: int, did: int, label: str = "NERc_04.1") -> tuple[int, int, int]:
    """One critical room: its marker on page 1 at (100, 100) and on page 2 at
    (800, 500). Returns (asset id, marker on page 1, marker on page 2)."""
    with SessionLocal() as db:
        e1 = Element(project_id=pid, drawing_id=did, page=1, kind="NER_CRITICAL",
                     label=label, x=100, y=100, origin="manual", confidence=1.0,
                     confirmed=True, total_ports=48)
        db.add(e1)
        db.commit()
        migrate.backfill_assets(db, db.get(Project, pid))
        db.refresh(e1)
        asset_id = e1.ner_asset_id
        e2 = Element(project_id=pid, drawing_id=did, page=2, kind="NER_CRITICAL",
                     label=label, x=800, y=500, origin="manual", confidence=1.0,
                     confirmed=True, ner_asset_id=asset_id, link_status="confirmed")
        db.add(e2)
        db.commit()
        return asset_id, e1.id, e2.id


def _teap_room(pid: int, did: int, code: str, page: int, x: float, y: float, ports: int = 4):
    with SessionLocal() as db:
        db.add(DrawingSpace(project_id=pid, drawing_id=did, page=page, space_code=code,
                            name="Office", x=x, y=y, confidence=0.9))
        tf = db.execute(select(TeapFile).where(TeapFile.project_id == pid)).scalars().first()
        db.add(TeapSpace(project_id=pid, teap_file_id=tf.id, space_key=f"ALA|{code}",
                         space_code=code, space_name="Office", total_ports=ports,
                         lines=[{"row": 1, "product": "Laptop", "code": "LAP", "qty": ports,
                                 "ports": ports, "rule": "A", "explanation": ""}]))
        db.commit()


def _label_at(doc, page: int, label: str) -> list[tuple[float, float]]:
    return [((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2) for r in doc[page - 1].search_for(label)]


# ---------------------------------------------------------------------------
# 1. A room on two sheets
# ---------------------------------------------------------------------------

def test_the_marked_pdf_draws_each_marker_where_its_own_sheet_draws_it(client):
    pid, did = _project("two-sheets-pdf")
    _room_on_two_sheets(pid, did)
    _teap_room(pid, did, "OFF.200", 2, 800, 540)
    assert client.post(f"/api/projects/{pid}/matching/run").status_code == 200
    with SessionLocal() as db:
        p = db.get(Project, pid)
        result, ners = S.analyse(db, p)
        S.save_scenario(db, p, "As-Is", "as_is", result, ners)
        db.commit()
    r = client.post(f"/api/projects/{pid}/reports/pdf?background=false")
    assert r.status_code == 200, r.text
    from app.config import settings
    doc = pymupdf.open(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    on1, on2 = _label_at(doc, 1, "NERc_04.1"), _label_at(doc, 2, "NERc_04.1")
    doc.close()
    assert len(on1) == 1 and abs(on1[0][0] - 100) < 30 and abs(on1[0][1] - 100) < 30, on1
    assert any(abs(x - 800) < 30 and abs(y - 500) < 30 for x, y in on2), \
        "the room is on page 2 too, where page 2 draws it"


def test_the_distance_is_measured_from_the_marker_on_the_spaces_sheet(client):
    pid, did = _project("two-sheets-measure")
    _room_on_two_sheets(pid, did)
    _teap_room(pid, did, "OFF.201", 2, 800, 530)
    assert client.post(f"/api/projects/{pid}/matching/run").status_code == 200
    with SessionLocal() as db:
        result, _ = S.analyse(db, db.get(Project, pid))
    a = next(a for a in result.assignments if a.space_key.endswith("OFF.201"))
    assert a.ner_label == "NERc_04.1"
    assert a.evaluation.effective_distance_m == pytest.approx(30.0)
    assert a.cross_level is False and a.ner_page == 2


def test_a_drag_on_another_sheet_leaves_the_canonical_point_alone(client):
    pid, did = _project("two-sheets-drag")
    asset_id, _m1, m2 = _room_on_two_sheets(pid, did)
    r = client.patch(f"/api/projects/{pid}/elements/{m2}", json={"x": 820, "y": 510})
    assert r.status_code == 200, r.text
    with SessionLocal() as db:
        a = db.get(NerAsset, asset_id)
        assert (a.canonical_page, a.canonical_x, a.canonical_y) == (1, 100, 100)


def test_a_drag_on_the_canonical_sheet_moves_the_canonical_point(client):
    pid, did = _project("two-sheets-drag-own")
    asset_id, m1, _m2 = _room_on_two_sheets(pid, did)
    assert client.patch(f"/api/projects/{pid}/elements/{m1}",
                        json={"x": 120, "y": 90}).status_code == 200
    with SessionLocal() as db:
        a = db.get(NerAsset, asset_id)
        assert (a.canonical_page, a.canonical_x, a.canonical_y) == (1, 120, 90)


def test_the_start_up_pass_puts_a_copied_canonical_point_back(client):
    """What the ALA database held: page 1, page 2's coordinates."""
    pid, did = _project("two-sheets-repair")
    asset_id, _m1, _m2 = _room_on_two_sheets(pid, did)
    with SessionLocal() as db:
        a = db.get(NerAsset, asset_id)
        a.canonical_x, a.canonical_y = 800, 500          # page stays 1
        db.commit()
        p = db.get(Project, pid)
        # Read right even before the repair …
        point, others = S.asset_points(a)
        assert (point.page, point.x, point.y) == (1, 100, 100)
        assert [(q.page, q.x, q.y) for q in others] == [(2, 800, 500)]
        # … and the stored value is put right, once, audited.
        assert migrate.repair_canonical_positions(db, p) == {"canonical_repaired": 1}
        db.refresh(a)
        assert (a.canonical_page, a.canonical_x, a.canonical_y) == (1, 100, 100)
        assert migrate.repair_canonical_positions(db, p) == {"canonical_repaired": 0}


def test_a_scenario_saved_before_2200_is_drawn_where_the_markers_are(client):
    pid, did = _project("two-sheets-old-snapshot")
    asset_id, _m1, _m2 = _room_on_two_sheets(pid, did)
    _teap_room(pid, did, "OFF.204", 1, 110, 120)
    with SessionLocal() as db:
        p = db.get(Project, pid)
        result, ners = S.analyse(db, p)
        sc = S.save_scenario(db, p, "As-Is", "as_is", result, ners)
        payload = dict(sc.result)
        # the shape a 2.19.0 snapshot had: one wrong point, no "other"
        payload["ners"] = [{k: v for k, v in n.items() if k != "other"} | {"x": 800, "y": 500}
                           for n in payload["ners"]]
        sc.result = payload
        db.commit()
        _, replay = S.scenario_view(db, p, "As-Is")
    room = next(n for n in replay if n.key == f"A{asset_id}")
    assert [(q.page, q.x, q.y) for q in room.points] == [(1, 100, 100), (2, 800, 500)]


# ---------------------------------------------------------------------------
# 2. An unconfirmed match is served by the nearest NER
# ---------------------------------------------------------------------------

def test_an_unconfirmed_match_is_served_and_is_not_an_exception():
    scale = Scale(metres_per_unit=1.0, label="1u=1m", source="manual", confirmed=True)
    area = Area(key="ALA|SPP.002", name="SPP Work Room", page=8,
                label_point=Point(100, 100, 8))
    res = compute_coverage(areas=[(area, 3.0, DemandType.GENERAL, None, False)],
                           ners=[NER(key="A1", label="NERc_10.1", kind="NER_CRITICAL",
                                     point=Point(120, 100, 8), critical=True)],
                           scale_for_page={8: scale}, rules=DistanceRules())
    a = res.assignments[0]
    assert a.ner_label == "NERc_10.1" and a.status is Status.SERVED
    assert a.match_confirmed is False and a.evaluation.requires_human_validation
    assert res.exceptions == [] and res.uncovered == []


# ---------------------------------------------------------------------------
# 3. The workbook
# ---------------------------------------------------------------------------

def test_the_ners_sheet_carries_full_name_space_code_and_notes(client):
    import openpyxl

    from app.config import settings
    pid, did = _project("workbook-notes")
    asset_id, _m1, _m2 = _room_on_two_sheets(pid, did)
    _teap_room(pid, did, "OFF.202", 1, 110, 120)
    r = client.patch(f"/api/projects/{pid}/ners/{asset_id}",
                     json={"values": {"space_code": "TEC.010-004",
                                      "notes": "Move to the athlete's medical station."}})
    assert r.status_code == 200, r.text
    assert r.json()["space_code"] == "TEC.010-004"
    assert client.post(f"/api/projects/{pid}/matching/run").status_code == 200
    with SessionLocal() as db:
        p = db.get(Project, pid)
        result, ners = S.analyse(db, p)
        S.save_scenario(db, p, "As-Is", "as_is", result, ners)
        db.commit()
    r = client.post(f"/api/projects/{pid}/reports/xlsx?background=false")
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    ws = wb["NERs and ports"]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["NER", "Full name", "Space code"]
    assert head[-1] == "Notes"
    row = [c.value for c in ws[2]]
    assert row[:3] == ["NERc_04.1", "ALA_NERc_04.1", "TEC.010-004"]
    assert row[-1] == "Move to the athlete's medical station."
    assert row[head.index("Utilisation %")] == '=IF(OR(I2="",I2=0),"",H2/I2)'


def test_the_sheets_come_in_his_order_with_his_tab_colours(client):
    import openpyxl

    from app.config import settings
    from app.domain.reports import SHEET_ORDER
    pid, did = _project("workbook-order")
    _room_on_two_sheets(pid, did)
    _teap_room(pid, did, "OFF.203", 1, 110, 120)
    with SessionLocal() as db:
        db.add(Element(project_id=pid, drawing_id=did, page=1, kind="WIRELESS",
                       label="WIFI-1", x=115, y=125, origin="manual", confidence=1.0,
                       confirmed=True))
        db.commit()
    assert client.post(f"/api/projects/{pid}/matching/run").status_code == 200
    with SessionLocal() as db:
        p = db.get(Project, pid)
        result, ners = S.analyse(db, p)
        S.save_scenario(db, p, "As-Is", "as_is", result, ners)
        db.commit()
    r = client.post(f"/api/projects/{pid}/reports/xlsx?background=false")
    wb = openpyxl.load_workbook(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    names = wb.sheetnames
    assert names == [n for n in SHEET_ORDER if n in names], names
    # 4.0.0 (Marcelo, 16 Oct 2026, from a screenshot of the tab bar)
    assert names[:5] == ["TEAP by NER", "NERs and ports", "Wireless points", "Summary",
                         "TEAP and coverage"]
    colour = {ws.title: (ws.sheet_properties.tabColor.rgb if ws.sheet_properties.tabColor
                         else None) for ws in wb.worksheets}
    for yellow in ("NERs and ports", "Wireless points"):
        assert colour[yellow].endswith("FFFF00"), (yellow, colour[yellow])
    assert colour["TEAP by NER"].endswith("FF0000")
    for plain in ("Summary", "TEAP and coverage", "Exceptions", "Scales"):
        assert colour[plain] is None, plain
    assert wb.active.title == "TEAP by NER"
