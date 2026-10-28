"""HTTP API for the NER register — the V2 screens.

Every mutating endpoint takes an ``actor`` and, where the decision is one an
engineer would be asked to defend, a ``reason``. Nothing here deletes a row.
"""

from __future__ import annotations

from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import ner_register as R
from .. import sheets as SH
from .. import services as S
from ..config import settings
from ..db import get_session
from ..models import (Criticality, Element, KeepDecision, NerAsset, NerConsolidation,
                      NerStatus, Project, SnapshotKind)

router = APIRouter()


def _project(db: Session, project_id: int) -> Project:
    p = db.get(Project, project_id)
    if p is None:
        raise HTTPException(404, "project not found")
    return p


def _same_names(db, p) -> None:
    """Refresh the "same name as another NER" validation (4.2.0)."""
    from .. import migrate as MG
    try:
        MG.flag_same_name_ners(db, p)
    except Exception:                                 # noqa: BLE001 — a hint, never a failure
        db.rollback()


def _guard(fn, *a, **kw):
    try:
        return fn(*a, **kw)
    except R.NameInUse as exc:                 # 409: the name is another NER's (4.2.0)
        raise HTTPException(409, str(exc)) from exc
    except R.AssetNotFound as exc:             # 404: no such NER here (4.8.0)
        raise HTTPException(404, str(exc)) from exc
    except R.RegisterError as exc:
        raise HTTPException(422, str(exc)) from exc


# ---------------------------------------------------------------------------
# Reading the register
# ---------------------------------------------------------------------------

@router.get("/projects/{project_id}/ners")
def list_ners(project_id: int, include_removed: bool = Query(False),
              db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    rows = R.list_assets(db, p, include_removed=include_removed)
    live = [r for r in rows if r["is_live"]]
    return {
        "ners": rows,
        "summary": {
            "total": len(rows),
            "live": len(live),
            "critical_confirmed": sum(1 for r in live
                                      if r["classification"] == Criticality.CRITICAL
                                      and r["classification_confirmed"]),
            "criticality_pending": sum(1 for r in live
                                       if not r["classification_confirmed"]),
            "suggested": sum(1 for r in rows if r["status"] == NerStatus.SUGGESTED),
            "rejected": sum(1 for r in rows if r["status"] == NerStatus.REJECTED),
            "consolidated": sum(1 for r in rows if r["merged_into_id"] is not None),
            "keep_undecided": sum(1 for r in live
                                  if r["keep_decision"] == KeepDecision.UNDECIDED),
            "capacity_unknown": sum(1 for r in live if not r["capacity"]["known"]),
            "occurrences": sum(r["occurrence_count"] for r in rows),
        },
        "statuses": list(NerStatus.ALL),
        "classifications": list(Criticality.ALL),
        "disclaimer": settings.disclaimer,
    }


@router.get("/projects/{project_id}/ners/{ner_id}")
def get_ner(project_id: int, ner_id: int, db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    return R.asset_dict(db, a, meta=R._page_meta(db, p.id), load=R._latest_load(db, p.id))


# ---------------------------------------------------------------------------
# Editing one NER
# ---------------------------------------------------------------------------

class RenameIn(BaseModel):
    name: str
    reason: str = ""
    actor: str = "user"            # ignored since 4.2.0: the signed-in person is the author
    #: Two NERs knowingly given one name (scripts only; the screens consolidate).
    same_name_ok: bool = False


@router.post("/projects/{project_id}/ners/{ner_id}/rename")
def rename_ner(project_id: int, ner_id: int, body: RenameIn,
               db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    before_id, before_code = a.id, a.code
    _guard(R.rename, db, p, a, body.name, actor=S.current_actor(), reason=body.reason,
           same_name_ok=body.same_name_ok)
    _same_names(db, p)
    return {"id": a.id, "code": a.code, "display_name": a.display_name,
            "previous_name": a.previous_name,
            "identity_unchanged": (a.id == before_id and a.code == before_code)}


# The bulk renamer ("apply {venue}-NER-{index} to every room") stood here
# until 2.7.0, and the naming convention with its suggested names and
# *Adopt* from 2.7.0 to 2.13.1. Both went (2.14.0) at Marcelo's request: every
# NER the app creates is named in sequence (NERc01 …, NER01 …) and renamed by
# hand, one row at a time, with the ordinary rename above. The event's naming
# structure is a reference panel in the interface and decides nothing.


class NextNameIn(BaseModel):
    critical: bool = False


@router.get("/projects/{project_id}/ner-next-name")
def next_name(project_id: int, critical: bool = False,
              db: Session = Depends(get_session)) -> dict:
    """The next free sequential name — what the viewer proposes for a NER
    drawn by hand."""
    p = _project(db, project_id)
    names = R.live_names(db, p.id)
    return {"name": R.next_sequential_name(names, critical),
            "critical": R.next_sequential_name(names, True),
            "non_critical": R.next_sequential_name(names, False)}


class ClassifyIn(BaseModel):
    classification: str
    confirmed: bool = False
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/classify")
def classify_ner(project_id: int, ner_id: int, body: ClassifyIn,
                 db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.classify, db, p, a, body.classification, confirmed=body.confirmed,
           actor=S.current_actor(), reason=body.reason)
    return {"id": a.id, "classification": a.classification,
            "classification_confirmed": a.classification_confirmed, "kind": a.kind}


class StatusIn(BaseModel):
    status: str
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/status")
def set_status(project_id: int, ner_id: int, body: StatusIn,
               db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.set_status, db, p, a, body.status, actor=S.current_actor(), reason=body.reason)
    return {"id": a.id, "status": a.status, "keep_decision": a.keep_decision}


class KeepIn(BaseModel):
    decision: str
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/keep")
def keep_ner(project_id: int, ner_id: int, body: KeepIn,
             db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.set_keep, db, p, a, body.decision, actor=S.current_actor(), reason=body.reason)
    return {"id": a.id, "keep_decision": a.keep_decision, "status": a.status}


class LockIn(BaseModel):
    locked: bool
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/lock")
def lock_ner(project_id: int, ner_id: int, body: LockIn,
             db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.set_lock, db, p, a, locked=body.locked, actor=S.current_actor(),
           reason=body.reason)
    return {"id": a.id, "locked": a.locked, "movable": a.movable}


class ConfirmIn(BaseModel):
    confirmed: bool = True
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/confirm")
def confirm_ner(project_id: int, ner_id: int, body: ConfirmIn,
                db: Session = Depends(get_session)) -> dict:
    """Sign a room off: its criticality and every marker of it on the sheet."""
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.set_confirmed, db, p, a, confirmed=body.confirmed, actor=S.current_actor(),
           reason=body.reason)
    return {"id": a.id, "confirmed": R.is_confirmed(a, list(a.occurrences)),
            "classification_confirmed": a.classification_confirmed, "kind": a.kind}


class BulkIn(BaseModel):
    action: str
    actor: str = "user"


@router.post("/projects/{project_id}/ners/bulk")
def bulk_ners(project_id: int, body: BulkIn,
              db: Session = Depends(get_session)) -> dict:
    """Lock all, unlock all, confirm all, unconfirm all — one audit row each."""
    p = _project(db, project_id)
    return _guard(R.bulk, db, p, body.action, actor=S.current_actor())


class FieldsIn(BaseModel):
    values: dict[str, Any]
    actor: str = "user"


@router.patch("/projects/{project_id}/ners/{ner_id}")
def patch_ner(project_id: int, ner_id: int, body: FieldsIn,
              db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.update_fields, db, p, a, body.values, actor=S.current_actor())
    return R.asset_dict(db, a, meta=R._page_meta(db, p.id), load=R._latest_load(db, p.id))


class MoveIn(BaseModel):
    page: int
    x: float
    y: float
    propose: bool = True
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/move")
def move_ner(project_id: int, ner_id: int, body: MoveIn,
             db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    # ``page`` is a sheet key (4.1.0) — the page number in a one-book project.
    try:
        did, page = SH.for_project(db, p.id).split(body.page)
    except SH.SheetError as e:
        raise HTTPException(404, str(e))
    _guard(R.move, db, p, a, page=page, x=body.x, y=body.y,
           propose=body.propose, actor=S.current_actor(), reason=body.reason, drawing_id=did)
    return R.asset_dict(db, a, meta=R._page_meta(db, p.id), load=R._latest_load(db, p.id))


class RemoveIn(BaseModel):
    reason: str
    actor: str = "user"


@router.post("/projects/{project_id}/ners/{ner_id}/remove")
def remove_ner(project_id: int, ner_id: int, body: RemoveIn,
               db: Session = Depends(get_session)) -> dict:
    """Controlled removal. The row and its history survive."""
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.soft_delete, db, p, a, reason=body.reason, actor=S.current_actor())
    _same_names(db, p)
    return {"id": a.id, "status": a.status, "deleted_at": a.deleted_at.isoformat(),
            "note": "the register keeps this row; nothing was deleted"}


@router.post("/projects/{project_id}/ners/{ner_id}/restore")
def restore_ner(project_id: int, ner_id: int, actor: str = "user",
                db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    a = _guard(R.get_asset, db, p, ner_id)
    _guard(R.restore, db, p, a, actor=S.current_actor())
    _same_names(db, p)
    return {"id": a.id, "status": a.status}


# ---------------------------------------------------------------------------
# Occurrences
# ---------------------------------------------------------------------------

class OccurrenceIn(BaseModel):
    asset_id: Optional[int] = None
    status: str = "confirmed"
    reason: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/occurrences/{element_id}/link")
def link_occurrence(project_id: int, element_id: int, body: OccurrenceIn,
                    db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    e = db.get(Element, element_id)
    if e is None or e.project_id != p.id:
        raise HTTPException(404, "occurrence not found in this project")
    _guard(R.set_occurrence_link, db, p, e, asset_id=body.asset_id,
           status=body.status, actor=S.current_actor(), reason=body.reason)
    return {"element_id": e.id, "ner_asset_id": e.ner_asset_id,
            "link_status": e.link_status}


@router.post("/projects/{project_id}/occurrences/{element_id}/detach")
def detach_occurrence(project_id: int, element_id: int, reason: str = "",
                      actor: str = "user",
                      db: Session = Depends(get_session)) -> dict:
    """Split one occurrence out into a NER of its own."""
    p = _project(db, project_id)
    e = db.get(Element, element_id)
    if e is None or e.project_id != p.id:
        raise HTTPException(404, "occurrence not found in this project")
    new = _guard(R.detach_occurrence, db, p, e, actor=S.current_actor(), reason=reason)
    return {"element_id": e.id, "new_ner_id": new.id, "code": new.code}


# ---------------------------------------------------------------------------
# Duplicates and consolidation
# ---------------------------------------------------------------------------

@router.get("/projects/{project_id}/ners/duplicates/candidates")
def duplicate_candidates(project_id: int, db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    rows = R.duplicate_candidates(db, p)
    return {
        "candidates": rows,
        "summary": {
            "pairs": len(rows),
            "auto_mergeable": sum(1 for r in rows if r["comparison"]["auto_mergeable"]),
            "needs_judgement": sum(1 for r in rows
                                   if not r["comparison"]["auto_mergeable"]),
        },
        "note": ("Nothing here has been consolidated. Each pair is a question with "
                 "its evidence and its differences; the tool never merges an "
                 "ambiguous pair on its own."),
    }


class ConsolidateIn(BaseModel):
    primary_id: int
    duplicate_id: int
    justification: str = ""
    force: bool = False
    actor: str = "user"


@router.post("/projects/{project_id}/ners/consolidate")
def consolidate(project_id: int, body: ConsolidateIn,
                db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    res = _guard(R.consolidate, db, p, body.primary_id, body.duplicate_id,
                  actor=S.current_actor(), justification=body.justification,
                  force=body.force)
    _same_names(db, p)
    return res


class SeparateIn(BaseModel):
    a_id: int
    b_id: int
    justification: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/ners/declare-separate")
def declare_separate(project_id: int, body: SeparateIn,
                     db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    return _guard(R.declare_separate, db, p, body.a_id, body.b_id,
                  actor=S.current_actor(), justification=body.justification)


@router.post("/projects/{project_id}/ner-consolidations/{consolidation_id}/undo")
def undo_consolidation(project_id: int, consolidation_id: int, reason: str = "",
                       actor: str = "user",
                       db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    res = _guard(R.undo_consolidation, db, p, consolidation_id, actor=S.current_actor(),
                  reason=reason)
    _same_names(db, p)
    return res


@router.get("/projects/{project_id}/ner-consolidations")
def consolidations(project_id: int, db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    return {"consolidations": R.consolidation_history(db, p)}


# ---------------------------------------------------------------------------
# Suggestions
# ---------------------------------------------------------------------------

# "Suggest NERs for uncovered areas" stood here. It ran the placement solver and
# wrote its new sites into the register as suggestions. Removed with the
# solver — see the note in api/routes.py — because a machine-chosen position on
# a plan it cannot read is not a suggestion a register should carry.

# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

# The register's own XLSX/CSV export stood here. It duplicated the analysis
# workbook, which carries every room and its ports on its own sheet, and nobody
# was reading two spreadsheets of the same rooms. The frozen-snapshot export
# below is a different thing and stays: it is what goes to the drawing team.

# ---------------------------------------------------------------------------
# Snapshots — change request and design freeze
# ---------------------------------------------------------------------------

class SnapshotIn(BaseModel):
    kind: str = "change_request"
    name: str
    reference: str = ""
    issued_to: str = ""
    note: str = ""
    actor: str = "user"


@router.post("/projects/{project_id}/register-snapshots", status_code=201)
def create_snapshot(project_id: int, body: SnapshotIn,
                    db: Session = Depends(get_session)) -> dict:
    """Freeze the register under a name. Immutable — there is no edit path."""
    p = _project(db, project_id)
    snap = _guard(R.create_snapshot, db, p, kind=body.kind, name=body.name,
                  reference=body.reference, issued_to=body.issued_to,
                  note=body.note, actor=S.current_actor())
    return R.snapshot_dict(snap, with_payload=True)


@router.get("/projects/{project_id}/register-snapshots")
def list_snapshots(project_id: int, db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    return {"snapshots": R.list_snapshots(db, p), "kinds": list(SnapshotKind.ALL)}


@router.get("/projects/{project_id}/register-snapshots/{snapshot_id}")
def get_snapshot(project_id: int, snapshot_id: int,
                 db: Session = Depends(get_session)) -> dict:
    p = _project(db, project_id)
    return R.snapshot_dict(_guard(R.get_snapshot, db, p, snapshot_id),
                           with_payload=True)


@router.get("/projects/{project_id}/register-snapshots/{snapshot_id}/export/{fmt}")
def export_snapshot(project_id: int, snapshot_id: int, fmt: str,
                    db: Session = Depends(get_session)):
    """The document that goes to the drawing team, or the frozen baseline."""
    import csv
    import io

    from fastapi.responses import Response

    p = _project(db, project_id)
    snap = _guard(R.get_snapshot, db, p, snapshot_id)
    items = (snap.payload or {}).get("change_items", [])
    register = (snap.payload or {}).get("register", [])
    keys = [k for k, _h in R.CHANGE_COLUMNS]
    headers = [h for _k, h in R.CHANGE_COLUMNS]
    stem = f"{snap.kind}-{snap.id}"

    def cell(item: dict, key: str):
        v = item.get(key, "")
        if isinstance(v, bool):
            return "yes" if v else "no"
        if isinstance(v, (list, tuple)):
            return "; ".join(str(x) for x in v)
        return v

    if fmt == "json":
        return R.snapshot_dict(snap, with_payload=True)

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, lineterminator="\n")
        w.writerow([snap.name])
        if snap.reference:
            w.writerow([f"Reference: {snap.reference}"])
        if snap.issued_to:
            w.writerow([f"Issued to: {snap.issued_to}"])
        w.writerow([f"Frozen: {snap.created_at.isoformat()} by {snap.created_by}"])
        w.writerow([])
        w.writerow(headers)
        for it in items:
            w.writerow([cell(it, k) for k in keys])
        w.writerow([])
        w.writerow([(snap.payload or {}).get("disclaimer", "")])
        return Response(buf.getvalue(), media_type="text/csv",
                        headers={"Content-Disposition":
                                 f'attachment; filename="{stem}.csv"'})

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Change request"
        ws["A1"] = snap.name
        ws["A1"].font = Font(bold=True, size=13)
        meta = [f"Reference: {snap.reference or '—'}",
                f"Issued to: {snap.issued_to or '—'}",
                f"Frozen: {snap.created_at.isoformat()} by {snap.created_by}",
                ("This list is frozen. It is the baseline the next drawing "
                 "version will be checked against — reissue rather than edit.")]
        for i, line in enumerate(meta, start=2):
            ws.cell(row=i, column=1, value=line)
        ws.append([])
        ws.append(headers)
        head_row = ws.max_row
        for c in ws[head_row]:
            c.font = Font(bold=True)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for it in items:
            ws.append([cell(it, k) for k in keys])
        ws.freeze_panes = f"A{head_row + 1}"
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=head_row, column=i).column_letter].width = \
                min(52, max(14, len(h) + 4))

        full = wb.create_sheet("Register at that moment")
        rkeys = [k for k, _h in R.REGISTER_COLUMNS]
        full.append([h for _k, h in R.REGISTER_COLUMNS])
        for c in full[1]:
            c.font = Font(bold=True)
        for row in register:
            full.append([R._join(row.get(k, "")) if isinstance(row.get(k), (list, tuple))
                         else row.get(k, "") for k in rkeys])
        full.freeze_panes = "C2"

        note = wb.create_sheet("Notice")
        note["A1"] = (snap.payload or {}).get("disclaimer", "")
        note["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        note.column_dimensions["A"].width = 110
        note["A3"] = ("Every row carries the internal code as well as the name. "
                      "The code is stable; names are renumbered between "
                      "revisions, so quote the code when there is any doubt "
                      "about which room is meant.")
        note["A3"].alignment = Alignment(wrap_text=True, vertical="top")

        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            buf.getvalue(),
            media_type=("application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"),
            headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'})

    raise HTTPException(422, "format must be one of: csv, xlsx, json")
