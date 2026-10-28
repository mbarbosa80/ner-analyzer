"""4.5.0 — indoor or outdoor for each Wi-Fi point.

Marcelo, 23 Oct 2026: beside each Wi-Fi, a way to say whether the service is
given indoors (a covered area) or outdoors (an open area), in a column next to
the Wi-Fi, as NERs already have — a detail that helps the LLD team understand
the setting each service is proposed for.

The field is said by a person and stored on the point; the API takes it for a
wireless point only; the workbook carries it beside "Type" on the Wireless
points sheet and counts it on the Summary, read off the points as they are
when exported. The screens are held by source checks, as the rest of this
suite does for the interface.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import openpyxl
import pymupdf
import pytest

TMP = tempfile.mkdtemp(prefix="ner-450-")
os.environ.setdefault("NER_DATA_DIR", TMP)
os.environ.setdefault("NER_DATABASE_URL", f"sqlite:///{TMP}/t450.sqlite3")

from fastapi.testclient import TestClient                          # noqa: E402

from app import migrate                                            # noqa: E402
from app import services as S                                      # noqa: E402
from app.db import SessionLocal, init_db                           # noqa: E402
from app.domain import reports as R                                # noqa: E402
from app.main import app                                           # noqa: E402
from app.models import (Drawing, DrawingSpace, Element, Page, Project,  # noqa: E402
                        TeapFile, TeapSpace)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "frontend" / "src"
DWISA = "Wireless Connection - Shared Internet Access (SIA) Wireless LAN"


def _src(rel: str) -> str:
    return (SRC / rel).read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def client():
    init_db()
    c = TestClient(app)
    from app import auth as A
    with SessionLocal() as db:
        try:
            A.create_user(db, username="tester", password="a-long-test-phrase",
                          display_name="Test User", is_admin=True, actor="tests")
        except A.AuthError:
            pass
    assert c.post("/api/auth/login", json={"username": "tester",
                                           "password": "a-long-test-phrase"}).status_code == 200
    return c


def _pdf(path: str) -> str:
    doc = pymupdf.open()
    doc.new_page(width=1000, height=700)
    doc.save(path)
    doc.close()
    return path


def _line(product, qty, ports, code="X", row=1):
    return {"row": row, "product": product, "code": code, "qty": qty, "ports": ports,
            "rule": "A", "explanation": ""}


@pytest.fixture(scope="module")
def project(client):
    """One sheet, two spaces with wireless, one NER, three Wi-Fi points."""
    with SessionLocal() as db:
        p = Project(name="wifi-env-450", venue_code="TST")
        db.add(p)
        db.flush()
        d = Drawing(project_id=p.id, filename="d.pdf", stored_path=_pdf(f"{TMP}/w450.pdf"),
                    sha256="x" * 64, page_count=1)
        db.add(d)
        db.flush()
        db.add(Page(drawing_id=d.id, number=1, width=1000, height=700, metres_per_unit=0.1,
                    scale_label="1u=0.1m", scale_source="manual", scale_confirmed=True,
                    level="LEVEL 01"))
        tf = TeapFile(project_id=p.id, filename="teap.xlsx", stored_path=f"{TMP}/t.xlsx",
                      sha256="y" * 64)
        db.add(tf)
        db.flush()
        for code, name, x in (("OFF.010", "Office", 100), ("PLZ.020", "Plaza", 600)):
            db.add(DrawingSpace(project_id=p.id, drawing_id=d.id, page=1, space_code=code,
                                name=name, x=x, y=100, confidence=0.9))
            db.add(TeapSpace(project_id=p.id, teap_file_id=tf.id, space_key=f"TST|{code}",
                             space_code=code, space_name=name, total_ports=2,
                             lines=[_line("Laptop", 1, 1, "LAP", 1),
                                    _line(DWISA, 1, 1, "DWISA", 2)],
                             wireless_required=True))
        db.add(Element(project_id=p.id, drawing_id=d.id, page=1, kind="NER_NONCRITICAL",
                       label="NER 1", x=300, y=160, origin="manual", confidence=1.0,
                       confirmed=True, total_ports=48))
        db.commit()
        pid = p.id
        migrate.backfill_assets(db, db.get(Project, pid))
        db.commit()
    assert client.post(f"/api/projects/{pid}/matching/run").status_code == 200
    ids = []
    for x in (110, 610, 800):
        r = client.post(f"/api/projects/{pid}/elements",
                        json={"page": 1, "kind": "WIRELESS", "label": "", "x": x, "y": 110})
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])
    return pid, ids


def _patch(client, pid, eid, body):
    return client.patch(f"/api/projects/{pid}/elements/{eid}", json=body)


# ---------------------------------------------------------------------------
# The field
# ---------------------------------------------------------------------------

def test_a_new_point_is_not_said(client, project):
    pid, ids = project
    rows = {e["id"]: e for e in client.get(f"/api/projects/{pid}/elements").json()}
    assert all(rows[i]["environment"] is None for i in ids)


def test_a_point_is_said_indoor_or_outdoor_and_cleared(client, project):
    pid, ids = project
    assert _patch(client, pid, ids[0], {"environment": "indoor"}).status_code == 200
    assert _patch(client, pid, ids[1], {"environment": "Outdoor "}).status_code == 200
    rows = {e["id"]: e for e in client.get(f"/api/projects/{pid}/elements").json()}
    assert rows[ids[0]]["environment"] == "indoor"
    assert rows[ids[1]]["environment"] == "outdoor", "trimmed and lower-cased"
    # "" is "not said" again
    assert _patch(client, pid, ids[2], {"environment": "outdoor"}).status_code == 200
    assert _patch(client, pid, ids[2], {"environment": ""}).status_code == 200
    rows = {e["id"]: e for e in client.get(f"/api/projects/{pid}/elements").json()}
    assert rows[ids[2]]["environment"] is None


def test_any_other_word_is_refused(client, project):
    pid, ids = project
    r = _patch(client, pid, ids[0], {"environment": "covered"})
    assert r.status_code == 422
    rows = {e["id"]: e for e in client.get(f"/api/projects/{pid}/elements").json()}
    assert rows[ids[0]]["environment"] == "indoor", "a refused value changes nothing"


def test_a_ner_marker_is_refused_here(client, project):
    pid, _ = project
    ner = next(e for e in client.get(f"/api/projects/{pid}/elements").json()
               if e["kind"] == "NER_NONCRITICAL")
    r = _patch(client, pid, ner["id"], {"environment": "outdoor"})
    assert r.status_code == 422 and "NERs screen" in r.text


def test_the_change_is_audited_and_leaves_the_analysis_current(client, project):
    pid, ids = project
    from sqlalchemy import select
    from app.models import AuditLog
    with SessionLocal() as db:
        rows = db.execute(select(AuditLog).where(
            AuditLog.project_id == pid, AuditLog.action == "wireless.environment",
            AuditLog.entity_id == ids[1])).scalars().all()
    assert any((r.payload or {}).get("to") == "outdoor" for r in rows)
    # Nothing the engine measures changed: the last run stays current.
    assert "wireless.environment" in S._ANALYSIS_NEUTRAL
    with SessionLocal() as db:
        p = db.get(Project, pid)
        result, ners = S.analyse(db, p)
        S.save_scenario(db, p, "As-Is", "as_is", result, ners)
        db.commit()
    assert _patch(client, pid, ids[1], {"environment": "indoor"}).status_code == 200
    assert _patch(client, pid, ids[1], {"environment": "outdoor"}).status_code == 200
    with SessionLocal() as db:
        assert S.analysis_stale(db, db.get(Project, pid)) is None


def test_the_wireless_points_list_carries_it(client, project):
    pid, ids = project
    rows = client.get(f"/api/projects/{pid}/spaces/assignments").json()["wireless_points"]
    by_id = {w["id"]: w for w in rows}
    assert by_id[ids[0]]["environment"] == "indoor"
    assert by_id[ids[1]]["environment"] == "outdoor"


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------

def test_the_sheet_has_the_column_beside_type():
    assert R.WIFI_SHEET_HEADERS[:3] == ["Wireless point", "Type", "Indoor / outdoor"]


def test_the_export_reads_the_points_as_they_are_now(client, project):
    pid, ids = project
    with SessionLocal() as db:
        p = db.get(Project, pid)
        result, ners = S.analyse(db, p)
        S.save_scenario(db, p, "As-Is", "as_is", result, ners)
        db.commit()
    # Said after the run: the export must still carry it.
    assert _patch(client, pid, ids[2], {"environment": "outdoor"}).status_code == 200
    r = client.post(f"/api/projects/{pid}/reports/xlsx?background=false")
    assert r.status_code == 200, r.text
    from app.config import settings
    wb = openpyxl.load_workbook(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    ws = wb["Wireless points"]
    head = [c.value for c in ws[1]]
    assert head == R.WIFI_SHEET_HEADERS
    col = head.index("Indoor / outdoor")
    labels = {e["id"]: e["label"] for e in client.get(f"/api/projects/{pid}/elements").json()}
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    # The first column names the point by its space when it has one; find the
    # rows by their position on the sheet instead: one per point, in order.
    values = [r[col] or "" for r in ws.iter_rows(min_row=2, values_only=True)]
    assert sorted(values) == ["indoor", "outdoor", "outdoor"], (values, rows, labels)
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Totals: Wi-Fi points indoor"] == 1
    assert summary["Totals: Wi-Fi points outdoor"] == 2
    assert summary["Totals: Wi-Fi points indoor/outdoor not said"] == 0


def test_without_wifi_info_the_run_value_is_used():
    # The run carries it on each point, and the exporter falls back on it
    # when it is not handed the points as they are now.
    src = (ROOT / "backend" / "app" / "domain" / "reports.py").read_text(encoding="utf-8")
    assert 'else a.area.meta.get("environment")) or ""' in src
    svc = (ROOT / "backend" / "app" / "services.py").read_text(encoding="utf-8")
    assert '"environment": e.environment})' in svc


# ---------------------------------------------------------------------------
# The screens
# ---------------------------------------------------------------------------

def test_demand_per_space_has_the_select_beside_type():
    s = _src("pages/panels.tsx")
    i = s.index('{t("Wireless point")}</sw2.Th>')
    head = s[i:i + 900]
    assert head.index('k="ap_kind"') < head.index('k="environment"') < head.index('k="page"')
    assert "setWirelessEnvironment(w, e.target.value)" in s
    assert "api.patchElement(p!.id, w.id, { environment: env })" in s


def test_the_inspector_has_the_select_for_a_wifi_point():
    v = _src("pages/Viewer.tsx")
    i = v.index("DEMAND_KINDS.has(selected.data.kind) && (() => {")
    block = v[i:i + 6000]
    assert "void patchSelected({ environment: ev.target.value })" in block
    for word in ('t("Indoor")', 't("Outdoor")', 't("— not said —")'):
        assert word in block


def test_the_new_text_is_in_portuguese():
    pt = _src("lib/strings.pt.ts")
    assert ('"Where the service is given: indoor (a covered area) or outdoor (an open area). '
            'A column of the Wireless points sheet in the workbook.":') in pt
