"""2.14.0 — the accommodation-matrix (AAM) space vocabulary, sequential NER names, one AP, the
indoor/outdoor field, the clean export.

Marcelo, 29 Sep 2026:

* *"Procurar primeiro pelas siglas, seguidas de número e usar nomes similares
  quando as siglas não estiverem presentes. Você começou a nomear muita coisa
  que antes não era nomeada … 'cable under the bridge', 'GEN load zone',
  'L145m'. Coisas que nada têm a ver com nomes de espaços."*
* *"Acabe com a sugestão de nomes, mantenha tudo como NERcXX e NERxx e crie
  tudo em sequência, que eu vou renomear tudo na mão mesmo."*
* *"Não diferenciar mais AP interno de externo, deixar tudo somente como AP."*
* *"Eu preciso dizer se ela é externa ou interna e colocar isso em uma coluna
  na planilha em Excel, na aba onde ficam as NERs."*
* *"Na exportação do PDF … limpar aquilo que não é de origem do arquivo."*

The PDFs here are synthetic — never a real venue document in a test.
"""

from __future__ import annotations

import os
import tempfile

import pymupdf
import pytest

TMP = tempfile.mkdtemp(prefix="ner-2140-")
os.environ.setdefault("NER_DATA_DIR", TMP)
os.environ.setdefault("NER_DATABASE_URL", f"sqlite:///{TMP}/t2140.sqlite3")

from sqlalchemy import select                                      # noqa: E402

from app import migrate                                            # noqa: E402
from app import ner_register as R                                  # noqa: E402
from app import services as S                                      # noqa: E402
from app.db import SessionLocal, init_db                           # noqa: E402
from app.domain import aam_vocab as V                              # noqa: E402
from app.domain.pdf_ingest import extract_drawing                  # noqa: E402
from app.domain.reports import strip_pdf_annotations               # noqa: E402
from app.models import (AuditLog, Drawing, DrawingSpace, Element, NerAsset,  # noqa: E402
                        Project, SpaceMatchRow, TeapSpace)
from tests.test_wireless_tie_250 import client  # noqa: E402,F401

# The list of known space names is the organiser's material, kept in
# backend/app/reference_private and left out of a public copy. Tests that
# need it skip when it is absent; the rest run either way.
NEEDS_PRIVATE_NAMES = pytest.mark.skipif(
    not V.aam_names(),
    reason="the private list of space names is not in this copy (backend/app/reference_private)")


@pytest.fixture(scope="module", autouse=True)
def _db():
    init_db()


# ---------------------------------------------------------------------------
# The vocabulary
# ---------------------------------------------------------------------------

@NEEDS_PRIVATE_NAMES
def test_the_bundled_aam_list_is_there_and_leaves_infrastructure_out():
    names = V.aam_names()
    assert len(names) > 250
    fas = {fa for fa, _ in names}
    assert {"ACR", "BRS", "PRS", "SPT", "TEC", "MED"} <= fas
    assert "VNI" not in fas and "MPS" not in fas, "cable bridges and partner footprints are not rooms"
    assert ("ACR", "Venue Accreditation Office (VAO)") in names


@pytest.mark.parametrize("label,expected", [
    pytest.param("Venue Accreditation Office (VAO)", True, marks=NEEDS_PRIVATE_NAMES),
    pytest.param("VENUE MEDIA CENTRE (VMC)", True, marks=NEEDS_PRIVATE_NAMES),
    pytest.param("L9m (31ft) Venue Media Centre (VMC)", True, marks=NEEDS_PRIVATE_NAMES),      # a length tag before the room
    pytest.param("Equipment control", True, marks=NEEDS_PRIVATE_NAMES),
    pytest.param("Athletes Locker Room - Women", True, marks=NEEDS_PRIVATE_NAMES),
    pytest.param("Broadcast Mixed Zone", True, marks=NEEDS_PRIVATE_NAMES),
    ("CABLE GOES UNDER BRIDGE", False),
    ("CABLE THROUGH PROPOSED PEDESTRIAN BRIDGE", False),
    ("L145m (476ft)", False),
    ("GEN Load Zone", False),
    ("STRM SEWR", False),
    ("1 / 410E1-CSL-O", False),
    ("KEY ENLARGED PLAN - LEVEL 01", False),
    ("TECHNOLOGY PATHWAY", False),
    ("UP", False),
    ("TMJ TMJ", False),
])
def test_a_label_is_a_space_only_when_it_reads_as_a_known_name(label, expected):
    assert V.SpaceVocabulary().is_space_name(label) is expected, label


def test_the_teap_names_extend_the_vocabulary():
    assert not V.SpaceVocabulary().is_space_name("Boat Storage Barn")
    assert V.SpaceVocabulary(["Boat Storage Barn"]).is_space_name("BOAT STORAGE BARN")


def test_the_length_tag_is_stripped_from_the_name():
    assert V.strip_length_tag("L9m (31ft) Venue Media Centre (VMC)") == "Venue Media Centre (VMC)"
    assert V.strip_length_tag("Venue Media Centre (VMC)") == "Venue Media Centre (VMC)"


def _sheet_without_codes(path, labels):
    doc = pymupdf.open()
    p = doc.new_page(width=1400, height=800)
    p.insert_text((1250, 700), "DRAWING TITLE", fontsize=7)
    p.insert_text((1250, 712), "FLOOR PLAN LEVEL 01 - AREA A1", fontsize=11)
    p.insert_text((1250, 730), "DRAWING NUMBER", fontsize=7)
    p.insert_text((1250, 742), "TST-411A1-XXX-O", fontsize=11)
    for i, text in enumerate(labels):
        p.insert_text((100 + (i % 4) * 300, 100 + (i // 4) * 120), text, fontsize=12)
    doc.save(str(path))
    doc.close()
    return path


@NEEDS_PRIVATE_NAMES
def test_the_extractor_keeps_the_rooms_and_drops_the_notes(tmp_path):
    book = _sheet_without_codes(tmp_path / "ctn.pdf", [
        "Venue Media Centre (VMC)", "CABLE GOES UNDER BRIDGE", "L18m (60ft)",
        "Athlete Medical Station", "GEN Load Zone", "STRM SEWR", "Equipment Control",
    ])
    ex = extract_drawing(str(book), known_names=[])
    assert sorted(s.name for s in ex.spaces) == [
        "Athlete Medical Station", "Equipment Control", "Venue Media Centre (VMC)"]
    assert all(s.space_code is None for s in ex.spaces)
    # without a vocabulary the old reading stands — every label
    assert len(extract_drawing(str(book)).spaces) == 7


def test_a_key_plan_yields_no_spaces_at_all(tmp_path):
    doc = pymupdf.open()
    p = doc.new_page(width=1400, height=800)
    p.insert_text((1250, 700), "DRAWING TITLE", fontsize=7)
    p.insert_text((1250, 712), "KEY ENLARGED PLAN - LEVEL 01", fontsize=11)
    p.insert_text((100, 100), "Venue Media Centre (VMC)", fontsize=12)
    p.insert_text((100, 200), "Office", fontsize=8)
    p.insert_text((100, 206), "OFF.050", fontsize=4)
    doc.save(str(tmp_path / "key.pdf"))
    ex = extract_drawing(str(tmp_path / "key.pdf"), known_names=[])
    assert ex.pages[0].key_plan and ex.spaces == []


# ---------------------------------------------------------------------------
# The start-up clean-up of an older book
# ---------------------------------------------------------------------------

@NEEDS_PRIVATE_NAMES
def test_the_cleanup_removes_untied_junk_and_keeps_tied_and_coded_spaces():
    with SessionLocal() as db:
        p = Project(name="clean-1", venue_code="TST")
        db.add(p)
        db.commit()
        d = Drawing(project_id=p.id, filename="old.pdf", stored_path="/nowhere/old.pdf",
                    page_count=2, extract_summary={"spaces": 5})
        db.add(d)
        db.flush()
        from app.models import Page
        db.add(Page(drawing_id=d.id, number=1, width=100, height=100, key_plan=True))
        db.add(Page(drawing_id=d.id, number=2, width=100, height=100, key_plan=False))
        rows = {}
        for key, (page, code, name) in {
            "coded": (2, "OFF.050", "Office"),
            "room": (2, None, "Venue Media Centre (VMC)"),
            "tagged": (2, None, "L9m (31ft) Venue Media Centre (VMC)"),
            "junk": (2, None, "CABLE GOES UNDER BRIDGE"),
            "junk_tied": (2, None, "GEN Load Zone"),
            "on_key": (1, "OFF.050", "Office"),
        }.items():
            sp = DrawingSpace(project_id=p.id, drawing_id=d.id, page=page, space_code=code,
                              name=name, x=10, y=10, origin="extracted")
            db.add(sp)
            db.flush()
            rows[key] = sp.id
        from app.models import TeapFile
        tf = TeapFile(project_id=p.id, filename="t.xlsx", stored_path="/nowhere/t.xlsx")
        db.add(tf)
        db.flush()
        t = TeapSpace(project_id=p.id, teap_file_id=tf.id, space_key="TST|X", space_code="X",
                      space_name="x", locked_drawing_space_id=rows["junk_tied"])
        db.add(t)
        db.commit()
        out = migrate.cleanup_spaces_2140(db, p)
        assert out == {"spaces_removed_not_a_name": 1, "spaces_removed_key_plan": 1,
                       "spaces_renamed": 1}
        left = {s.id: s for s in db.execute(select(DrawingSpace).where(
            DrawingSpace.project_id == p.id)).scalars()}
        assert set(left) == {rows["coded"], rows["room"], rows["tagged"], rows["junk_tied"]}
        assert left[rows["tagged"]].name == "Venue Media Centre (VMC)"
        assert db.get(Drawing, d.id).extract_summary["spaces_pass"] == S.SPACES_PASS
        # once per book
        assert migrate.cleanup_spaces_2140(db, p) == {
            "spaces_removed_not_a_name": 0, "spaces_removed_key_plan": 0, "spaces_renamed": 0}
        log = db.execute(select(AuditLog).where(AuditLog.project_id == p.id,
                                                AuditLog.action == "drawing.spaces_cleaned")).scalar_one()
        assert "p2 CABLE GOES UNDER BRIDGE" in log.payload["names"]


# ---------------------------------------------------------------------------
# Sequential names
# ---------------------------------------------------------------------------

def test_the_next_name_follows_the_highest_of_its_series():
    assert R.next_sequential_name([], True) == "NERc_01.1"
    assert R.next_sequential_name([], False) == "NER_20.1"   # 2.18.0: non-critical from 20 (v2.0)
    assert R.next_sequential_name(["NERc01", "NERc02", "NER01"], True) == "NERc_03.1"
    assert R.next_sequential_name(["NERc01", "NERc02", "NER21"], False) == "NER_22.1"
    # spelling aside: "NERc 07", "ner-3", "NER 21", "NERc_07.1", "VHH_NERc_07.2" all count;
    # a hand name does not
    assert R.next_sequential_name(["NERc 07", "Critical NER 12", "VEE-NER-01"], True) == "NERc_08.1"
    assert R.next_sequential_name(["VHH_NERc_07.2", "NERc_03.1"], True) == "NERc_08.1"
    assert R.next_sequential_name(["ner-3", "NER 21"], False) == "NER_22.1"
    # a number below the series' start never pulls it down (2.15.0)
    assert R.next_sequential_name(["NER 02"], False) == "NER_20.1"
    # a collision with any existing name is skipped
    assert R.next_sequential_name(["NER 20.1", "NERc01"], False) == "NER_21.1"


def test_a_marker_the_extractor_read_is_named_in_sequence_and_a_hand_one_keeps_its_name(client):
    with SessionLocal() as db:
        p = Project(name="seq-1", venue_code="TST")
        db.add(p)
        db.commit()
        d = Drawing(project_id=p.id, filename="b.pdf", stored_path="/nowhere/b.pdf", page_count=1)
        db.add(d)
        db.flush()
        from app.models import Page
        db.add(Page(drawing_id=d.id, number=1, width=1000, height=1000, level="LEVEL 01"))
        db.add(Page(drawing_id=d.id, number=2, width=1000, height=1000, level="LEVEL 01"))
        for label, kind, method, origin, page in (
                ("Critical NER 8", "NER_CRITICAL", "pdf_extraction", "extracted", 1),
                ("NER 19", "NER_NONCRITICAL", "pdf_extraction", "extracted", 1),
                ("NER 19", "NER_NONCRITICAL", "pdf_extraction", "extracted", 2),   # drawn on two sheets
                ("My rack", "NER_NONCRITICAL", "manual", "manual", 1)):
            db.add(Element(project_id=p.id, drawing_id=d.id, page=page, kind=kind, label=label,
                           x=100, y=100, origin=origin, detection_method=method,
                           confidence=0.9, evidence=[]))
        db.commit()
        out = migrate.backfill_assets(db, p)
        assert out["created"] == 3 and out["auto_merged"] == 1
        rows = {a.display_name: a for a in db.execute(select(NerAsset).where(
            NerAsset.project_id == p.id)).scalars()}
        assert set(rows) == {"NERc_01.1", "NER_20.1", "My rack"}
        assert rows["NERc_01.1"].suggested_name == "Critical NER 8"
        assert rows["NER_20.1"].suggested_name == "NER 19" and "NER 19" in rows["NER_20.1"].aliases
        assert rows["My rack"].suggested_name == "My rack"
        # one name: both markers of NER 19 are labelled NER_20.1, the print kept
        both = [e for e in rows["NER_20.1"].occurrences]
        assert len(both) == 2 and all(e.label == "NER_20.1" and e.nearby_text == "NER 19" for e in both)
        pid = p.id
    # the register's next name follows
    nxt = client.get(f"/api/projects/{pid}/ner-next-name").json()
    assert (nxt["critical"], nxt["non_critical"]) == ("NERc_02.1", "NER_21.1")
    # the naming endpoints of 2.7.0–2.13.1 are gone
    assert client.get(f"/api/projects/{pid}/naming").status_code == 404


# ---------------------------------------------------------------------------
# Indoor / outdoor
# ---------------------------------------------------------------------------

def test_indoor_outdoor_is_a_register_field_with_a_workbook_column(client):
    import openpyxl

    from tests.test_wireless_tie_250 import _analyse, _match, _ner, _project, _room, _teap, OFFICE
    pid = _project("env-1")
    _room(pid, "OFF.060", "Office", 100, 100)
    _teap(pid, "OFF.060", "Office", OFFICE)
    eid = _ner(pid, "NER A", 100, 120, total_ports=24)
    aid = next(e for e in client.get(f"/api/projects/{pid}/elements").json() if e["id"] == eid)["ner_asset_id"]
    _match(client, pid)
    r = client.patch(f"/api/projects/{pid}/ners/{aid}", json={"values": {"environment": "outdoor"}})
    assert r.status_code == 200, r.text
    assert client.get(f"/api/projects/{pid}/ners/{aid}").json()["environment"] == "outdoor"
    assert client.patch(f"/api/projects/{pid}/ners/{aid}",
                        json={"values": {"environment": "roof"}}).status_code == 422
    # the viewer's marker carries it too
    assert next(e for e in client.get(f"/api/projects/{pid}/elements").json()
                if e["id"] == eid)["environment"] == "outdoor"
    _analyse(pid)
    r = client.post(f"/api/projects/{pid}/reports/xlsx?background=false")
    assert r.status_code == 200, r.text
    from app.config import settings
    wb = openpyxl.load_workbook(settings.exports_dir / f"project_{pid}" / r.json()["file"])
    ws = wb["NERs and ports"]
    head = [c.value for c in ws[1]]
    assert head[5] == "Indoor / outdoor"
    assert [c.value for c in ws[2]][5] == "outdoor"
    summary = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value
               for ws in (wb["Summary"],) for i in range(2, wb["Summary"].max_row + 1)}
    assert summary["Totals: NERs outdoor"] == 1 and summary["Totals: NERs indoor"] == 0


# ---------------------------------------------------------------------------
# The clean export
# ---------------------------------------------------------------------------

def test_the_export_takes_the_reviewers_annotations_off_the_copy(tmp_path):
    doc = pymupdf.open()
    p = doc.new_page(width=800, height=600)
    p.insert_text((100, 100), "Office", fontsize=12)          # page content: stays
    a = p.add_freetext_annot(pymupdf.Rect(300, 300, 400, 330), "NER 14", fontsize=12)
    a.update()
    ink = p.add_ink_annot([[(500, 500), (520, 510), (540, 500)]])
    ink.update()
    src = tmp_path / "annotated.pdf"
    doc.save(str(src))
    doc.close()
    d = pymupdf.open(str(src))
    assert sum(1 for _ in d[0].annots()) == 2
    assert strip_pdf_annotations(d) == 2
    assert sum(1 for _ in d[0].annots()) == 0
    assert "Office" in d[0].get_text()
    d.close()
    # the file on disk is untouched
    assert sum(1 for _ in pymupdf.open(str(src))[0].annots()) == 2
